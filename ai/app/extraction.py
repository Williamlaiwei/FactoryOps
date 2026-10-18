"""Bounded document extraction. Content is reference data, never instructions."""
import csv
from dataclasses import dataclass
import io
from pathlib import PurePath
import re
from urllib.parse import urlparse
import zipfile

import openpyxl
import pdfplumber

MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 2_000_000
MAX_CHUNKS = 2000


class DocumentError(ValueError):
    pass


@dataclass
class Chunk:
    content: str
    metadata: dict


def safe_url(value: str) -> str:
    if not value:
        return ""
    parsed = urlparse(value)
    if len(value) > 500 or parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
        raise DocumentError("Source URL must be an HTTPS URL without credentials.")
    return value


def chunk_text(text: str, metadata: dict, size: int = 1200, overlap: int = 150):
    text = text.replace("\x00", "").strip()
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = text.rfind(" ", start + size // 2, end)
            if boundary > start:
                end = boundary
        value = text[start:end].strip()
        if value:
            yield Chunk(value, dict(metadata, char_start=start, char_end=end))
        if end == len(text):
            break
        start = max(start + 1, end - overlap)


def extract(filename: str, data: bytes):
    if not data or len(data) > MAX_BYTES:
        raise DocumentError("Choose a nonempty PDF, CSV or XLSX file no larger than 20 MB.")
    suffix = PurePath(filename.lower()).suffix
    segments = []
    summary = {"format": suffix.lstrip("."), "warnings": []}
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise DocumentError("The file is not a valid PDF.")
        try:
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                if len(pdf.pages) > 200:
                    raise DocumentError("PDF documents may contain at most 200 pages.")
                summary["pages"] = len(pdf.pages)
                for number, page in enumerate(pdf.pages, 1):
                    text = page.extract_text() or ""
                    if text.strip():
                        segments.append((text, {"page": number, "kind": "text"}))
                    else:
                        summary["warnings"].append(f"Page {number} has no extractable text; OCR may be required.")
                    for index, table in enumerate(page.extract_tables(), 1):
                        text = "\n".join(" | ".join(str(cell or "") for cell in row) for row in table)
                        if text.strip():
                            segments.append((text, {"page": number, "kind": "table", "table": index}))
        except DocumentError:
            raise
        except Exception as error:
            raise DocumentError("PDF could not be read. Use an unencrypted, text-based PDF.") from error
    elif suffix == ".csv":
        try:
            rows = csv.reader(io.StringIO(data.decode("utf-8-sig")), strict=True)
            headers = next(rows, [])
            if not headers or len(headers) > 100:
                raise DocumentError("CSV needs a header and at most 100 columns.")
            count = 0
            for row in rows:
                if not any(cell.strip() for cell in row):
                    continue
                count += 1
                if count > 20000 or len(row) != len(headers):
                    raise DocumentError("CSV has inconsistent columns or more than 20,000 rows.")
                values = dict(zip(headers, row))
                text = "; ".join(f"{key}: {value}" for key, value in values.items() if value)
                metadata = {"kind": "row", "row_start": rows.line_num, "row_end": rows.line_num}
                if values.get("SourceUrl"):
                    metadata["source_url"] = safe_url(values["SourceUrl"])
                segments.append((text, metadata))
            summary["rows"] = count
        except (UnicodeError, csv.Error) as error:
            raise DocumentError("CSV could not be read. Use UTF-8 with consistent columns and correctly quoted values.") from error
    elif suffix == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 2000 or sum(entry.file_size for entry in archive.infolist()) > 80 * 1024 * 1024:
                    raise DocumentError("Workbook expands beyond the supported size.")
            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                if len(book.worksheets) > 20:
                    raise DocumentError("Workbooks may contain at most 20 sheets.")
                count = 0
                for sheet in book.worksheets:
                    for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if len(row) > 100:
                            raise DocumentError("Worksheets may contain at most 100 columns.")
                        if not any(value is not None for value in row):
                            continue
                        count += 1
                        if count > 20000:
                            raise DocumentError("Workbooks may contain at most 20,000 nonempty rows.")
                        text = " | ".join(str(value) if value is not None else "" for value in row)
                        segments.append((text, {"kind": "row", "sheet": sheet.title, "row_start": index, "row_end": index}))
                summary["sheets"] = book.sheetnames
            finally:
                book.close()
        except DocumentError:
            raise
        except Exception as error:
            raise DocumentError("Workbook could not be read. Use an unencrypted XLSX file.") from error
    else:
        raise DocumentError("Supported knowledge formats are PDF, CSV and XLSX.")
    if sum(len(text) for text, _ in segments) > MAX_TEXT:
        raise DocumentError("Document contains too much extracted text; split it into smaller files.")
    chunks = [chunk for text, metadata in segments for chunk in chunk_text(text, metadata)]
    if not chunks:
        raise DocumentError("No text was extracted. Scanned PDFs need OCR before upload.")
    if len(chunks) > MAX_CHUNKS:
        raise DocumentError("Document creates too many chunks; split it into smaller files.")
    summary["chunks"] = len(chunks)
    return chunks, summary
